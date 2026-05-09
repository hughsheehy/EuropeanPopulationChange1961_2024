import csv
import json
from pathlib import Path
from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parent
ARDECO_CSV = BASE_DIR / "ARDECO_Local_Population_Time-Series–1961-2024.csv"
GISCO_UK_2011 = BASE_DIR / "LAU_RG_01M_2011_4326.geojson"
UK_WARD_XLSX = BASE_DIR / "6697_localunit_variable_details_2013.xlsx"
UK_WARD_GEOJSON = BASE_DIR / "WD_MAY_2025_UK_BGC_V2_-6160789691130864750.geojson"
OUTPUT_CSV = BASE_DIR / "unmatched_uk_crosswalk.csv"


def load_unmatched_uk_codes() -> list[str]:
    with ARDECO_CSV.open("r", encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.DictReader(fh))

    uk_csv = {row["gisco_id"] for row in rows if row["gisco_id"].startswith("UK")}

    with GISCO_UK_2011.open("r", encoding="utf-8") as fh:
        geo = json.load(fh)

    gisco_uk = {
        feature["properties"].get("GISCO_ID")
        for feature in geo["features"]
        if str(feature["properties"].get("GISCO_ID", "")).startswith("UK")
    }

    normalized_csv = {f"UK_{code[2:]}" for code in uk_csv}
    unmatched = sorted(normalized_csv - gisco_uk)
    return unmatched


def load_legacy_lookup(unmatched_suffixes: set[str]) -> dict[str, dict[str, str]]:
    wb = load_workbook(UK_WARD_XLSX, read_only=True, data_only=True)
    ws = wb["main_codistward"]

    lookup: dict[str, dict[str, str]] = {}
    for row in ws.iter_rows(min_row=3, values_only=True):
        code, name, _, district, county, region, country = row[:7]
        if code in unmatched_suffixes:
            lookup[code] = {
                "legacy_code": code or "",
                "legacy_name": name or "",
                "legacy_district": district or "",
                "legacy_county": county or "",
                "legacy_region": region or "",
                "legacy_country": country or "",
            }
    return lookup


def normalize(text: str) -> str:
    return " ".join((text or "").strip().lower().replace("&", "and").split())


def load_modern_wards():
    with UK_WARD_GEOJSON.open("r", encoding="utf-8") as fh:
        geo = json.load(fh)

    records = []
    for feature in geo["features"]:
        props = feature["properties"]
        records.append(
            {
                "wd25cd": props.get("WD25CD", ""),
                "wd25nm": props.get("WD25NM", ""),
                "lad25cd": props.get("LAD25CD", ""),
                "lad25nm": props.get("LAD25NM", ""),
                "norm_name": normalize(props.get("WD25NM", "")),
                "norm_lad": normalize(props.get("LAD25NM", "")),
            }
        )
    return records


def choose_match(legacy: dict[str, str], wards: list[dict[str, str]]) -> dict[str, str]:
    name = normalize(legacy["legacy_name"])
    district = normalize(legacy["legacy_district"])

    exact_pair = [w for w in wards if w["norm_name"] == name and w["norm_lad"] == district]
    if len(exact_pair) == 1:
        w = exact_pair[0]
        return {
            "match_type": "exact_name_and_district",
            "wd25cd": w["wd25cd"],
            "wd25nm": w["wd25nm"],
            "lad25cd": w["lad25cd"],
            "lad25nm": w["lad25nm"],
            "candidate_count": "1",
        }
    if len(exact_pair) > 1:
        w = exact_pair[0]
        return {
            "match_type": "exact_name_and_district_multiple",
            "wd25cd": w["wd25cd"],
            "wd25nm": w["wd25nm"],
            "lad25cd": w["lad25cd"],
            "lad25nm": w["lad25nm"],
            "candidate_count": str(len(exact_pair)),
        }

    exact_name = [w for w in wards if w["norm_name"] == name]
    if len(exact_name) == 1:
        w = exact_name[0]
        return {
            "match_type": "name_only",
            "wd25cd": w["wd25cd"],
            "wd25nm": w["wd25nm"],
            "lad25cd": w["lad25cd"],
            "lad25nm": w["lad25nm"],
            "candidate_count": "1",
        }
    if len(exact_name) > 1:
        w = exact_name[0]
        return {
            "match_type": "name_only_multiple",
            "wd25cd": w["wd25cd"],
            "wd25nm": w["wd25nm"],
            "lad25cd": w["lad25cd"],
            "lad25nm": w["lad25nm"],
            "candidate_count": str(len(exact_name)),
        }

    same_district = [w for w in wards if w["norm_lad"] == district]
    if same_district:
        return {
            "match_type": "district_only",
            "wd25cd": "",
            "wd25nm": "",
            "lad25cd": same_district[0]["lad25cd"],
            "lad25nm": same_district[0]["lad25nm"],
            "candidate_count": str(len(same_district)),
        }

    return {
        "match_type": "unmatched",
        "wd25cd": "",
        "wd25nm": "",
        "lad25cd": "",
        "lad25nm": "",
        "candidate_count": "0",
    }


def main():
    unmatched = load_unmatched_uk_codes()
    suffixes = {code[3:] for code in unmatched}
    legacy_lookup = load_legacy_lookup(suffixes)
    wards = load_modern_wards()

    fieldnames = [
        "ardeco_gisco_id",
        "legacy_code",
        "legacy_name",
        "legacy_district",
        "legacy_county",
        "legacy_region",
        "legacy_country",
        "match_type",
        "candidate_count",
        "wd25cd",
        "wd25nm",
        "lad25cd",
        "lad25nm",
    ]

    rows = []
    for normalized_code in unmatched:
        suffix = normalized_code[3:]
        legacy = legacy_lookup.get(
            suffix,
            {
                "legacy_code": suffix,
                "legacy_name": "",
                "legacy_district": "",
                "legacy_county": "",
                "legacy_region": "",
                "legacy_country": "",
            },
        )
        match = choose_match(legacy, wards) if legacy["legacy_name"] else {
            "match_type": "missing_in_legacy_lookup",
            "candidate_count": "0",
            "wd25cd": "",
            "wd25nm": "",
            "lad25cd": "",
            "lad25nm": "",
        }

        row = {
            "ardeco_gisco_id": f"UK{suffix}",
            **legacy,
            **match,
        }
        rows.append(row)

    with OUTPUT_CSV.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)

    print(f"Wrote {len(rows)} rows to {OUTPUT_CSV}")


if __name__ == "__main__":
    main()
